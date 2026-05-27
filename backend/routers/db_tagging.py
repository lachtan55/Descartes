"""
/api/db  — document text extraction and AI tagging endpoints.

POST /api/db/extract-text   { file_url, filename }  →  { text, char_count, extraction_method }
POST /api/db/suggest-tags   { text, filename }       →  { tags, suggested_title, summary }
"""

from __future__ import annotations

import asyncio
import io
import json
import re
from pathlib import Path
from typing import Literal

import httpx
from fastapi import APIRouter, HTTPException
from pydantic import BaseModel

from backend import config
from backend.gemini import call_gemini_json

# ── Optional anthropic import ─────────────────────────────────────────────────
# Mirrors the blpapi guard pattern: safe to run without the package installed;
# the /suggest-tags route returns HTTP 503 with install instructions if missing.
try:
    import anthropic as _anthropic
    _HAS_ANTHROPIC = True
except ImportError:
    _anthropic = None  # type: ignore[assignment]
    _HAS_ANTHROPIC = False

router = APIRouter(prefix="/api/db", tags=["db-tagging"])

# ── Taxonomy (mirrors frontend taxonomy.ts) ───────────────────────────────────

FIXED_VALUES = {
    "market":     ["comdty_market", "financial_market", "stock_market", "crypto_market",
                   "energy_market", "realestate_market"],
    "instrument": ["stock", "ETF", "index", "futures", "spot", "options"],
    "comdty":     ["comdty_rare-metals", "comdty_colored-metals", "comdty_energy",
                   "comdty_soft", "comdty_grain", "comdty_hardware", "comdty_special"],
    "exchange":   ["HL_EXCHANGE", "IB_GETTEX", "IB_NASDAQ", "IB_NYSE", "IB_XETR"],
    "data_type":  ["price_data", "thesis", "research_report", "macro_data", "earnings",
                   "transcript", "news", "regulatory", "technical_analysis", "notes"],
    "special":    ["war_geopolitics", "short_squeeze", "insider_buy", "insider_sell",
                   "activist", "m_and_a", "earnings_beat", "earnings_miss", "bankruptcy",
                   "ipo", "spinoff", "supply_disruption"],
    "timeframe":  ["daily", "weekly", "monthly", "quarterly", "annual", "intraday"],
}

FREE_FORM_CATEGORIES = {"ticker", "currency", "macro"}

ALL_CATEGORIES = list(FIXED_VALUES.keys()) + list(FREE_FORM_CATEGORIES)

# ── Pydantic models ───────────────────────────────────────────────────────────

class ExtractRequest(BaseModel):
    file_url: str
    filename: str


class ExtractResponse(BaseModel):
    text: str
    char_count: int
    extraction_method: str


class SuggestRequest(BaseModel):
    text: str
    filename: str
    provider: Literal["gemini", "claude"] = "gemini"


class TagSuggestion(BaseModel):
    category: str
    value: str
    confidence: float


class SuggestResponse(BaseModel):
    tags: list[TagSuggestion]
    suggested_title: str
    summary: str


# ── Text extraction ───────────────────────────────────────────────────────────

async def _download_file(url: str) -> bytes:
    """Download a file via HTTP and return raw bytes."""
    async with httpx.AsyncClient(timeout=60.0, follow_redirects=True) as client:
        resp = await client.get(url)
        if resp.status_code != 200:
            raise HTTPException(
                status_code=502,
                detail=f"Failed to download file (HTTP {resp.status_code}): {url}",
            )
        return resp.content


def _extract_pdf(data: bytes) -> tuple[str, str]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages: list[str] = []
        failed = 0
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                # Some pages have malformed IndirectObjects in their geometry;
                # skip the bad page rather than failing the whole document.
                pages.append("")
                failed += 1
        text = "\n".join(pages)
        method = "pypdf" if failed == 0 else f"pypdf-partial ({len(reader.pages) - failed}/{len(reader.pages)} pages)"
        return text, method
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"PDF extraction failed: {e}")


def _extract_docx(data: bytes) -> tuple[str, str]:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs]
        return "\n".join(paragraphs), "python-docx"
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"DOCX extraction failed: {e}")


def _extract_xlsx(data: bytes) -> tuple[str, str]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    lines.append("\t".join(cells))
        return "\n".join(lines), "openpyxl"
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"XLSX extraction failed: {e}")


def _extract_csv(data: bytes) -> tuple[str, str]:
    try:
        text = data.decode("utf-8", errors="replace")
        return text, "csv-plain"
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"CSV extraction failed: {e}")


def _extract_text_bytes(data: bytes) -> tuple[str, str]:
    return data.decode("utf-8", errors="replace"), "plaintext"


def extract_text_from_bytes(filename: str, data: bytes) -> tuple[str, str]:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "pdf":
        return _extract_pdf(data)
    elif ext in ("docx", "doc"):
        return _extract_docx(data)
    elif ext == "xlsx":
        return _extract_xlsx(data)
    elif ext == "csv":
        return _extract_csv(data)
    else:
        # txt, md, or anything else → plain text
        return _extract_text_bytes(data)


# ── Shared tagging prompt ────────────────────────────────────────────────────

_TAXONOMY_DESCRIPTION: str = (
    "FIXED categories — use ONLY the listed values:\n"
    + json.dumps(FIXED_VALUES, indent=2)
    + "\n\nFREE-FORM categories — invent appropriate values:\n"
      "  ticker:    tickers mentioned (e.g. 'AAPL', 'GC', 'BTC-USD')\n"
      "  currency:  currencies referenced (e.g. 'USD', 'EUR', 'CNY')\n"
      "  macro:     macroeconomic themes (e.g. 'inflation', 'rate_hike', 'supply_chain')"
)


def _build_tagging_prompt(filename: str, text: str) -> str:
    return (
        "You are a financial research document tagger.\n\n"
        f"{_TAXONOMY_DESCRIPTION}\n\n"
        "Assign confidence 0.0–1.0 per tag; apply all relevant tags.\n\n"
        f"Filename: {filename}\n"
        f'Text (first 8000 chars):\n"""\n{text[:8000]}\n"""\n\n'
        "Return ONLY valid JSON — no markdown, no explanation:\n"
        '{\n'
        '  "suggested_title": "<concise title, max 80 chars>",\n'
        '  "summary": "<2-3 sentence summary>",\n'
        '  "tags": [{"category": "<cat>", "value": "<val>", "confidence": 0.0}]\n'
        '}'
    )


def _strip_fences(s: str) -> str:
    """Remove markdown code fences that some models add despite instructions."""
    s = s.strip()
    if s.startswith("```"):
        s = re.sub(r"^```[a-z]*\n?", "", s)
        s = re.sub(r"\n?```$", "", s)
    return s.strip()


# ── Route handlers ────────────────────────────────────────────────────────────

@router.post("/extract-text", response_model=ExtractResponse)
async def extract_text(req: ExtractRequest):
    """Download the file from PocketBase and extract its text content."""
    data = await _download_file(req.file_url)
    # Run synchronous parser (pypdf / python-docx / openpyxl) in a thread so
    # the event loop stays free for health-checks and other requests.
    text, method = await asyncio.to_thread(extract_text_from_bytes, req.filename, data)

    # Trim excessively long text to keep downstream calls manageable
    MAX_CHARS = 50_000
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]

    return ExtractResponse(
        text=text,
        char_count=len(text),
        extraction_method=method,
    )


@router.post("/suggest-tags", response_model=SuggestResponse)
async def suggest_tags(req: SuggestRequest):
    """Call the selected AI provider to suggest tags, title, and summary."""
    prompt = _build_tagging_prompt(req.filename, req.text)

    if req.provider == "claude":
        if not config.ANTHROPIC_API_KEY:
            raise HTTPException(
                status_code=400,
                detail="ANTHROPIC_API_KEY not configured",
            )
        if not _HAS_ANTHROPIC:
            raise HTTPException(
                status_code=503,
                detail=(
                    "anthropic package is not installed in the active Python environment. "
                    "Run: .\\venv\\Scripts\\python.exe -m pip install anthropic "
                    "then restart the FastAPI server."
                ),
            )
        try:
            # Run the blocking SDK call in a thread — keeps the event loop free.
            def _call_claude() -> dict:
                client = _anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY)
                msg = client.messages.create(
                    model="claude-haiku-4-5-20251001",
                    max_tokens=1024,
                    messages=[{"role": "user", "content": prompt}],
                )
                return json.loads(_strip_fences(msg.content[0].text))

            result: dict = await asyncio.to_thread(_call_claude)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=502, detail=f"Claude error: {exc}")
    else:
        # Gemini branch — call_gemini_json uses time.sleep; run in thread.
        result = await asyncio.to_thread(call_gemini_json, prompt)

    # ── Normalise and validate the response ───────────────────────────────────
    suggested_title = str(result.get("suggested_title", req.filename))[:80]
    summary = str(result.get("summary", ""))
    raw_tags = result.get("tags", [])

    tags: list[TagSuggestion] = []
    for t in raw_tags:
        if not isinstance(t, dict):
            continue
        category = str(t.get("category", "")).strip()
        value = str(t.get("value", "")).strip()
        try:
            confidence = float(t.get("confidence", 0.5))
        except (TypeError, ValueError):
            confidence = 0.5

        if category not in ALL_CATEGORIES:
            continue
        if category in FIXED_VALUES and value not in FIXED_VALUES[category]:
            continue
        if not value:
            continue

        tags.append(TagSuggestion(
            category=category,
            value=value,
            confidence=min(1.0, max(0.0, confidence)),
        ))

    return SuggestResponse(
        tags=tags,
        suggested_title=suggested_title,
        summary=summary,
    )
